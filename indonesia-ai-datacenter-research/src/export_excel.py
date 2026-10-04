"""Create the editable, source-linked analyst workbook with openpyxl.

Excel formulas own the workbook calculations. Initial formula caches are
independently calculated in Python so readers without a calculation engine can
inspect the saved results. This is not a claim of native Excel recalculation.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import textwrap
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

from config import ROOT, PROCESSED, assumptions, read_dataset, write_json
from scenario_model import evaluate

SHEETS = [
    'Executive Summary', 'Data Center Projects', 'Indonesia Electricity',
    'Electricity Tariffs', 'Water Tariffs', 'Climate', 'PUE', 'WUE',
    'AI Hardware', 'Scenario Model', 'Sensitivity', '2030 Outlook',
    'National Impact', 'Sources', 'Assumptions',
]
RAW_SHEETS = {
    'Data Center Projects': 'datacenter_projects',
    'Indonesia Electricity': 'indonesia_electricity',
    'Electricity Tariffs': 'electricity_tariffs', 'Water Tariffs': 'water_tariffs',
    'Climate': 'climate', 'PUE': 'pue_benchmarks', 'WUE': 'wue_benchmarks',
    'AI Hardware': 'ai_hardware',
}
BLUE, GREEN, NAVY, TEAL = '0000FF', '008000', '18354B', '326D78'
NUM = '#,##0.00;(#,##0.00);"–"'
INT = '#,##0;(#,##0);"–"'
PCT = '0.00%;(0.00%);"–"'
CASE_COL = {'low': 'C', 'base': 'D', 'high': 'E'}
MODEL_HEADERS = [
    'Case', 'IT capacity (MW)', 'Electrical utilization', 'PUE (facility/IT)',
    'WUE (L/IT kWh)', 'Electricity (Rp/kWh)', 'Water (Rp/m³)', 'Hours/year',
    'IT electricity (MWh/year)', 'Facility electricity (MWh/year)',
    'Facility electricity (TWh/year)', 'Average facility load (MW)',
    'Full IT load facility power (MW)', 'Direct water (m³/year)',
    'Direct water (m³/day)', 'Electricity cost (Rp/year)',
    'Water cost (Rp/year)', 'Total resource cost (Rp/year)',
    'Electricity cost share', 'Water cost share', 'Data class', 'Scope / notes',
]


def _scalar(value):
    if value is None or pd.isna(value):
        return None
    if hasattr(value, 'item'):
        value = value.item()
    if isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        try:
            return datetime.strptime(value, '%Y-%m-%d').date()
        except ValueError:
            pass
    return value


class Builder:
    def __init__(self, settings):
        self.a = settings
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        for name in SHEETS:
            ws = self.wb.create_sheet(name)
            ws.sheet_view.showGridLines = False
            ws.sheet_view.zoomScale = 85
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.sheet_properties.tabColor = NAVY if name == 'Executive Summary' else TEAL
        self.caches = {name: {} for name in SHEETS}
        self.tables = {}
        self.scenario_rows = {}
        self.outlook_rows = {}
        self.sensitivity_refs = {}
        self.wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)

    def put(self, sheet, address, value, fmt=None, editable=False):
        cell = self.wb[sheet][address]
        cell.value = _scalar(value)
        cell.font = Font(name='Arial', size=10, color=BLUE if editable else NAVY)
        cell.alignment = Alignment(vertical='center', horizontal='right' if isinstance(cell.value, (int, float)) else 'left')
        if fmt:
            cell.number_format = fmt
        elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
            cell.number_format = NUM
        elif isinstance(cell.value, (date, datetime)):
            cell.number_format = 'yyyy-mm-dd'
        if editable:
            cell.fill = PatternFill('solid', fgColor='FFF4D8')
        return cell

    def formula(self, sheet, address, formula, cache, fmt=NUM):
        cell = self.put(sheet, address, formula, fmt)
        cell.font = Font(name='Arial', size=10, color=GREEN if '!' in formula else '000000')
        cell.alignment = Alignment(horizontal='right', vertical='center')
        self.caches[sheet][address] = _scalar(cache)

    def link(self, sheet, address, source, cache, fmt=NUM):
        self.formula(sheet, address, f'=IF(ISNUMBER({source}),{source},"Missing input")', cache, fmt)

    def guarded(self, sheet, address, expression, inputs, cache, fmt=NUM):
        self.formula(sheet, address, f'=IF(COUNT({",".join(inputs)})={len(inputs)},{expression},"Missing input")', cache, fmt)

    def header(self, sheet, row, values):
        ws = self.wb[sheet]
        for col, value in enumerate(values, 1):
            cell = self.put(sheet, f'{get_column_letter(col)}{row}', value)
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=NAVY)
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
        ws.row_dimensions[row].height = 44

    def title(self, sheet, title, note, width=8):
        ws = self.wb[sheet]
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        self.put(sheet, 'A2', title).font = Font(name='Arial', size=17, bold=True, color=NAVY)
        ws.row_dimensions[2].height = 28
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width)
        self.put(sheet, 'A3', note).alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[3].height = 35
        ws.freeze_panes = 'C6'

    def raw_table(self, sheet, frame):
        ws = self.wb[sheet]
        self.tables[sheet] = frame
        self.header(sheet, 1, list(frame.columns))
        for r, values in enumerate(frame.itertuples(index=False, name=None), 2):
            for c, val in enumerate(values, 1):
                name = str(frame.columns[c - 1])
                cell = self.put(sheet, f'{get_column_letter(c)}{r}', val)
                if name == 'year' or name.endswith('_year') or name == 'month':
                    cell.number_format = '0'
                if name == 'url' and isinstance(val, str):
                    cell.hyperlink = val
                    cell.font = Font(name='Arial', size=10, color=GREEN, underline='single')
                if isinstance(cell.value, str):
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[r].height = 45
        for c, name in enumerate(frame.columns, 1):
            width = 19
            if name in ['notes', 'rationale', 'title']:
                width = 62
            elif name in ['facility', 'hardware', 'customer_category', 'typical_server_configuration']:
                width = 34
            elif name in ['url', 'scope']:
                width = 50
            elif name in ['source_id', 'data_class', 'month', 'year']:
                width = 14
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = '1:1'
        ws.print_area = ws.dimensions

    def fit_wrapped_rows(self, sheet):
        """Fit populated wrapped cells so source caveats are readable in Excel."""
        ws = self.wb[sheet]
        for row in ws.iter_rows():
            height = ws.row_dimensions[row[0].row].height or 16
            for cell in row:
                if not cell.alignment.wrap_text or not isinstance(cell.value, str):
                    continue
                # Presentation merges have their own deliberately sized note areas.
                if any(cell.coordinate in merged for merged in ws.merged_cells.ranges):
                    continue
                width = ws.column_dimensions[cell.column_letter].width or 13
                wrap_width = max(8, int(width * .86))
                lines = sum(max(1, len(textwrap.wrap(line, width=wrap_width))) for line in cell.value.split('\n'))
                height = max(height, 15 * lines + 8)
            ws.row_dimensions[row[0].row].height = min(409, height)

    def assumptions(self):
        a, s = self.a, 'Assumptions'
        self.title(s, 'Indonesia data center model assumptions',
                   'Edit blue cells. All three cases calculate simultaneously. A = reported, E = derived, B = benchmark, S = scenario. WUE uses IT electricity.', 8)
        self.header(s, 6, ['Expansion driver', 'Unit', 'Low', 'Base', 'High', 'Source IDs', 'Class', 'Definition / rationale'])
        for r, key, label, unit in [(7, 'utilization', 'Electrical utilization', 'fraction'), (8, 'pue', 'Expansion PUE', 'facility/IT energy'), (9, 'wue_l_per_kwh', 'Direct cooling WUE', 'L/IT kWh')]:
            self.put(s, f'A{r}', label)
            self.put(s, f'B{r}', unit)
            for case, col in CASE_COL.items():
                self.put(s, f'{col}{r}', a['scenarios'][case][key], PCT if key == 'utilization' else NUM, True)
            self.put(s, f'F{r}', ';'.join(a['assumptions_source_ids']))
            self.put(s, f'G{r}', 'S')
            self.put(s, f'H{r}', ' | '.join(f'{case}: {a["scenarios"][case]["rationale"]}' for case in CASE_COL)).alignment = Alignment(wrap_text=True, vertical='top')
            self.wb[s].row_dimensions[r].height = 72
        self.header(s, 11, ['Current fleet driver', 'Unit', 'Low', 'Base', 'High', 'Source IDs', 'Class', 'Definition / rationale'])
        for r, key, label, unit in [(12, 'utilization', 'Electrical utilization', 'fraction'), (13, 'pue', 'Current fleet PUE', 'facility/IT energy'), (14, 'wue_l_per_kwh', 'Direct cooling WUE', 'L/IT kWh')]:
            self.put(s, f'A{r}', label)
            self.put(s, f'B{r}', unit)
            for case, col in CASE_COL.items():
                self.put(s, f'{col}{r}', a['current_resource_scenarios'][case][key], PCT if key == 'utilization' else NUM, True)
            self.put(s, f'F{r}', 'WB008;WB011;WB017')
            self.put(s, f'G{r}', 'S')
            self.put(s, f'H{r}', 'Current fleet stress cases. PUE is not a measured capacity-weighted Indonesian fleet average.').alignment = Alignment(wrap_text=True)
            self.wb[s].row_dimensions[r].height = 44
        self.header(s, 16, ['Shared input', 'Value', 'Unit', 'Source ID', 'Class', 'Definition / rationale'])
        constants = [
            (17, 'Hours/year', a['hours_per_year'], 'hours', 'MODEL001', 'S', 'Non-leap annual steady-state run rate.'),
            (18, 'Electricity tariff', a['tariffs']['electricity_rp_per_kwh'], 'Rp/kWh', a['tariffs']['electricity_source_id'], 'A/S', a['tariffs']['electricity_rationale']),
            (19, 'Water tariff proxy', a['tariffs']['water_rp_per_m3'], 'Rp/m³', a['tariffs']['water_source_id'], 'A/S', a['tariffs']['water_rationale']),
            (20, 'Connection power factor', a['power_factor']['base'], 'fraction', 'MODEL001', 'S', a['power_factor']['rationale']),
            (21, 'Generation capacity factor', a['generation_capacity_factor'], 'fraction', 'MODEL001', 'S', a['generation_capacity_factor_rationale']),
            (22, 'Greater Jakarta baseline IT', a['baseline']['it_capacity_mw'], 'MW', a['baseline']['source_id'], 'E', a['baseline']['scope']),
            (23, 'Baseline observation date', a['baseline']['observation_date'], 'date', a['baseline']['source_id'], 'E', 'Dated market estimate. National electricity and water use are not metered here.'),
            (24, 'National denominator year', a['national_denominator']['year'], 'year', a['national_denominator']['source_id'], 'A', 'Consumption and PLN sales remain separate denominators.'),
            (25, 'National electricity consumption', self.den['electricity_consumption_twh'], 'TWh', self.den['source_id'], 'A', 'Links to the matching dated national source row. Not generation.'),
            (26, 'PLN electricity sales', self.den['pln_sales_twh'], 'TWh', self.den['source_id'], 'A', 'Includes Batam as defined in HEESI. Not all national electricity consumption.'),
            (27, 'Future consumption growth', a['national_denominator']['outlook_consumption_cagr'], 'fraction/year', 'MODEL001', 'S', a['national_denominator']['rationale']),
            (28, 'Sensitivity IT capacity', a['sensitivity_capacity_mw'], 'MW', 'MODEL001', 'S', 'One-driver tests hold the remaining base assumptions fixed.'),
            (29, 'Partial disclosed operational IT', self.inventory_mw, 'MW', 'Data Center Projects', 'E', 'Sum of included explicitly IT operational capacity. Partial inventory. Do not add to the market baseline.'),
        ]
        for r, label, val, unit, source, cls, note in constants:
            for c, v in enumerate([label, val, unit, source, cls, note], 1):
                self.put(s, f'{get_column_letter(c)}{r}', v, editable=c == 2 and r not in [25, 26, 29])
            self.wb[s][f'F{r}'].alignment = Alignment(wrap_text=True, vertical='top')
            self.wb[s].row_dimensions[r].height = 45
        for r, col in [(25, 'electricity_consumption_twh'), (26, 'pln_sales_twh')]:
            idx = self.tables['Indonesia Electricity'].columns.get_loc(col) + 1
            end = len(self.tables['Indonesia Electricity']) + 1
            years = f"'Indonesia Electricity'!$A$2:$A${end}"
            values = f"'Indonesia Electricity'!${get_column_letter(idx)}$2:${get_column_letter(idx)}${end}"
            lookup = f'INDEX({values},MATCH($B$24,{years},0))'
            expression = f'=IF(COUNTIFS({years},$B$24)=1,IF(ISBLANK({lookup}),"Missing input",IF(ISNUMBER({lookup}),{lookup},"Missing input")),"Missing input")'
            self.formula(s, f'B{r}', expression, self.den[col])
        p = self.tables['Data Center Projects']
        end = len(p) + 1
        contribution_col = get_column_letter(len(p.columns)+1)
        self.formula(s, 'B29', f"=SUM('Data Center Projects'!${contribution_col}$2:${contribution_col}${end})", self.inventory_mw)
        for r in [20, 21, 27]:
            self.wb[s][f'B{r}'].number_format = PCT
        self.wb[s]['B24'].number_format = '0'
        self.wb[s]['B17'].number_format = INT
        self.header(s, 31, ['Capacity case', 'IT capacity (MW)', 'Class', 'Source', 'Definition'])
        for r, mw in enumerate(a['capacity_scenarios_mw'], 32):
            for c, val in enumerate([f'{mw:g} MW IT', mw, 'S', 'MODEL001', 'Nameplate IT capacity, not a grid connection or measured consumption.'], 1):
                self.put(s, f'{get_column_letter(c)}{r}', val, editable=c == 2)
        self.header(s, 39, ['Outlook case', 'Year', 'IT capacity (MW)', 'Resource case', 'Source IDs', 'Capacity rationale'])
        row = 40
        for name, cfg in a['outlook'].items():
            for year, mw in cfg['capacity_mw'].items():
                for c, val in enumerate([name, int(year), mw, cfg['resource_case'], cfg['anchor_source_id'], cfg['rationale']], 1):
                    self.put(s, f'{get_column_letter(c)}{row}', val, '0' if c == 2 else None, editable=c == 3)
                self.wb[s][f'F{row}'].alignment = Alignment(wrap_text=True, vertical='top')
                self.wb[s].row_dimensions[row].height = 90
                self.outlook_rows[(name, int(year))] = row
                row += 1
        self.header(s, 50, ['Sensitivity driver', 'Test value', 'Source', 'Class', 'Definition'])
        row = 51
        for driver, values in a['sensitivity'].items():
            for value in values:
                self.put(s, f'A{row}', driver)
                self.put(s, f'B{row}', value, PCT if 'multiplier' in driver or driver == 'utilization' else NUM, True)
                self.put(s, f'C{row}', 'MODEL001')
                self.put(s, f'D{row}', 'S')
                self.put(s, f'E{row}', 'Independent sensitivity, not a joint engineering cooling curve.')
                self.sensitivity_refs[(driver, value)] = row
                row += 1
        self.header(s, 74, ['Official PLN forecast', 'Sales (TWh)', 'Year', 'Source ID', 'Class', 'Definition'])
        self.pln_forecast_rows = {}
        forecast = read_dataset('outlook_2030_2032')
        for r, year in enumerate([2030,2031,2032],75):
            match = forecast[forecast.year.eq(year)]
            value = float(match.iloc[0]['official_pln_sales_forecast_twh'])
            for c, val in enumerate(['RUPTL PLN-only forecast',value,year,'ELEC003','S','Official forecast of PLN sales. Excludes non-PLN demand.'],1):
                self.put(s,f'{get_column_letter(c)}{r}',val,'0' if c == 3 else None)
            self.pln_forecast_rows[year] = (r,value)
        ws = self.wb[s]
        for col, width in {'A': 35, 'B': 25, 'C': 21, 'D': 24, 'E': 24, 'F': 65, 'G': 14, 'H': 80}.items():
            ws.column_dimensions[col].width = width
        for rng, lower, upper in [('C7:E7', 0, 1), ('C12:E12', 0, 1), ('B20:B21', 0.001, 1), ('C8:E8', 1, 10), ('C13:E13', 1, 10), ('C9:E9', 0, 100), ('C14:E14', 0, 100)]:
            dv = DataValidation(type='decimal', operator='between', formula1=lower, formula2=upper)
            dv.error, dv.errorTitle, dv.showErrorMessage = 'Enter a numeric value in the displayed physical range.', 'Invalid input', True
            ws.add_data_validation(dv)
            dv.add(rng)
        ws.freeze_panes = 'C7'
        ws.print_area = 'A1:H77'

    def case_row(self, sheet, row, case, capacity_ref, capacity, current=False, scope='Expansion scenario', cls='S', driver_overrides=None):
        from copy import deepcopy
        a = deepcopy(self.a)
        if current:
            a['scenarios'] = a['current_resource_scenarios']
        if driver_overrides:
            for driver, value in driver_overrides.items():
                if driver in ['electricity_rp_per_kwh', 'water_rp_per_m3']:
                    a['tariffs'][driver] = value
                else:
                    a['scenarios'][case][driver] = value
        out = evaluate(capacity, case, a)
        self.put(sheet, f'A{row}', case.title())
        self.link(sheet, f'B{row}', capacity_ref, capacity)
        start = 12 if current else 7
        for col, offset, key in [('C', 0, 'utilization'), ('D', 1, 'pue'), ('E', 2, 'wue_l_per_kwh')]:
            self.link(sheet, f'{col}{row}', f"'Assumptions'!{CASE_COL[case]}{start+offset}", out[key], PCT if col == 'C' else NUM)
        self.link(sheet, f'F{row}', "'Assumptions'!$B$18", out['electricity_tariff_rp_per_kwh'])
        self.link(sheet, f'G{row}', "'Assumptions'!$B$19", out['water_tariff_rp_per_m3'])
        self.link(sheet, f'H{row}', "'Assumptions'!$B$17", self.a['hours_per_year'], INT)
        formulas = [
            ('I', f'B{row}*C{row}*H{row}', ['B','C','H'], 'it_electricity_mwh', NUM),
            ('J', f'I{row}*D{row}', ['I','D'], 'facility_electricity_mwh', NUM),
            ('K', f'J{row}/1000000', ['J'], 'facility_electricity_twh', NUM),
            ('L', f'J{row}/H{row}', ['J','H'], 'average_facility_load_mw', NUM),
            ('M', f'B{row}*D{row}', ['B','D'], 'full_it_load_facility_mw', NUM),
            ('N', f'I{row}*1000*E{row}/1000', ['I','E'], 'water_m3_per_year', INT),
            ('O', f'N{row}/365', ['N'], 'water_m3_per_day', INT),
            ('P', f'J{row}*1000*F{row}', ['J','F'], 'electricity_cost_rp', INT),
            ('Q', f'N{row}*G{row}', ['N','G'], 'water_cost_rp', INT),
            ('R', f'SUM(P{row}:Q{row})', ['P','Q'], 'total_resource_cost_rp', INT),
            ('S', f'IF(R{row}=0,"n.a.",P{row}/R{row})', ['P','R'], 'electricity_cost_share_pct', PCT),
            ('T', f'IF(R{row}=0,"n.a.",Q{row}/R{row})', ['Q','R'], 'water_cost_share_pct', PCT),
        ]
        for col, expression, inputs, key, fmt in formulas:
            cache = out[key] / 100 if key.endswith('_share_pct') and out[key] is not None else out[key]
            self.guarded(sheet, f'{col}{row}', expression, [f'{c}{row}' for c in inputs], cache, fmt)
        self.put(sheet, f'U{row}', cls)
        self.put(sheet, f'V{row}', scope).alignment = Alignment(wrap_text=True, vertical='top')
        self.wb[sheet].row_dimensions[row].height = 30
        return out

    def model(self):
        s = 'Scenario Model'
        self.title(s, 'Annual electricity, water and resource costs',
                   'IT MW × electrical utilization × annual hours gives IT MWh. PUE gives facility energy. Direct consumptive water uses WUE per IT kWh. Costs cover energy and the Jakarta water proxy.', 11)
        self.header(s, 5, MODEL_HEADERS)
        row = 6
        for cap_index, mw in enumerate(self.a['capacity_scenarios_mw'], 32):
            for case in CASE_COL:
                self.case_row(s, row, case, f"'Assumptions'!$B${cap_index}", mw)
                self.scenario_rows[(mw, case)] = row
                row += 1
        self.wb[s].auto_filter.ref = 'A5:V23'
        self.put(s, 'A27', 'Dated Greater Jakarta resource estimate').font = Font(name='Arial', size=13, bold=True, color=NAVY)
        self.header(s, 29, MODEL_HEADERS)
        for row, case in enumerate(CASE_COL, 30):
            self.case_row(s, row, case, "'Assumptions'!$B$22", self.a['baseline']['it_capacity_mw'], True, self.a['baseline']['scope'], 'E')
        self.put(s, 'A35', 'Partial disclosed operational IT inventory').font = Font(name='Arial', size=13, bold=True, color=NAVY)
        self.header(s, 36, MODEL_HEADERS)
        for row, case in enumerate(CASE_COL, 37):
            self.case_row(s, row, case, "'Assumptions'!$B$29", self.inventory_mw, True, 'Partial disclosed inventory. Not Indonesia total. Do not add to the overlapping Greater Jakarta market estimate.', 'E')
        self.size_model(s)

    def size_model(self, sheet):
        ws = self.wb[sheet]
        for col in range(1, 23):
            ws.column_dimensions[get_column_letter(col)].width = 24 if col in [16,17,18] else 20
        ws.column_dimensions['A'].width = 16
        ws.column_dimensions['V'].width = 72
        ws.print_title_rows = '5:5'
        ws.print_area = ws.dimensions

    def sensitivity(self):
        s = 'Sensitivity'
        self.title(s, 'Independent driver sensitivities at 1 GW IT',
                   'Each row changes one driver. Other drivers remain linked to the expansion base case. PUE and WUE combinations are independent stress tests, not validated cooling designs.', 11)
        self.header(s, 5, MODEL_HEADERS + ['Changed driver', 'Test value'])
        from copy import deepcopy
        row = 6
        for driver, values in self.a['sensitivity'].items():
            for value in values:
                override = {}
                if driver.endswith('_multiplier'):
                    key = 'electricity_rp_per_kwh' if driver.startswith('electricity') else 'water_rp_per_m3'
                    override[key] = self.a['tariffs'][key] * value
                else:
                    override[driver] = value
                out = self.case_row(s, row, 'base', "'Assumptions'!$B$28", self.a['sensitivity_capacity_mw'], driver_overrides=override)
                assumption_row = self.sensitivity_refs[(driver, value)]
                src = f"'Assumptions'!$B${assumption_row}"
                target = {'utilization':'C', 'pue':'D', 'wue_l_per_kwh':'E', 'electricity_tariff_multiplier':'F', 'water_tariff_multiplier':'G'}[driver]
                if driver.endswith('_multiplier'):
                    base_row = 18 if target == 'F' else 19
                    self.guarded(s, f'{target}{row}', f"'Assumptions'!$B${base_row}*{src}", [f"'Assumptions'!$B${base_row}",src], out['electricity_tariff_rp_per_kwh' if target == 'F' else 'water_tariff_rp_per_m3'])
                else:
                    self.link(s, f'{target}{row}', src, value, PCT if driver == 'utilization' else NUM)
                self.put(s, f'W{row}', driver)
                self.link(s, f'X{row}', src, value)
                row += 1
        self.wb[s].auto_filter.ref = f'A5:X{row-1}'
        matrix_row = row + 4
        self.put(s, f'A{matrix_row}', 'Joint cooling stress: total annual resource cost (Rp trillion)').font = Font(name='Arial', size=13, bold=True, color=NAVY)
        self.header(s, matrix_row + 2, ['PUE / WUE (L/IT kWh)'] + list(self.a['sensitivity']['wue_l_per_kwh']))
        for c, wue in enumerate(self.a['sensitivity']['wue_l_per_kwh'], 2):
            self.link(s, f'{get_column_letter(c)}{matrix_row+2}', f"'Assumptions'!B{self.sensitivity_refs[('wue_l_per_kwh',wue)]}", wue)
        for r, pue in enumerate(self.a['sensitivity']['pue'], matrix_row + 3):
            self.link(s, f'A{r}', f"'Assumptions'!B{self.sensitivity_refs[('pue',pue)]}", pue)
            for c, wue in enumerate(self.a['sensitivity']['wue_l_per_kwh'], 2):
                b = deepcopy(self.a)
                b['scenarios']['base'].update(pue=pue,wue_l_per_kwh=wue)
                out = evaluate(self.a['sensitivity_capacity_mw'],'base',b)
                inputs = ["'Assumptions'!$B$28", "'Assumptions'!$D$7", "'Assumptions'!$B$17", f'$A{r}', f'{get_column_letter(c)}${matrix_row+2}', "'Assumptions'!$B$18", "'Assumptions'!$B$19"]
                expr = f"'Assumptions'!$B$28*'Assumptions'!$D$7*'Assumptions'!$B$17*($A{r}*1000*'Assumptions'!$B$18+{get_column_letter(c)}${matrix_row+2}*'Assumptions'!$B$19)/1000000000000"
                self.guarded(s, f'{get_column_letter(c)}{r}',expr,inputs,out['total_resource_cost_rp']/1e12)
        self.size_model(s)
        self.wb[s].column_dimensions['W'].width = 33
        self.wb[s].column_dimensions['X'].width = 20

    def outlook(self):
        s = '2030 Outlook'
        self.title(s, 'Source-anchored 2030–2032 capacity and resource outlook',
                   'Incomplete nationwide documented capacity combines the Greater Jakarta market anchor with separately sourced Surabaya, Batam and Bintan. Analyst realization assumptions. Unknown capacity is excluded.', 11)
        self.header(s, 5, MODEL_HEADERS + ['Outlook case','Year','National consumption scenario (TWh)','Share of future national consumption','Capacity anchor source IDs','Official PLN sales forecast (TWh)','Share of official PLN forecast','PLN forecast source ID'])
        row = 6
        for (name, year), ar in self.outlook_rows.items():
            cfg = self.a['outlook'][name]
            out = self.case_row(s,row,cfg['resource_case'],f"'Assumptions'!C{ar}",cfg['capacity_mw'][year],scope=cfg['rationale'])
            self.put(s,f'W{row}',name)
            self.formula(s,f'X{row}',f"='Assumptions'!B{ar}",year,'0')
            den = self.den['electricity_consumption_twh']*(1+self.a['national_denominator']['outlook_consumption_cagr'])**(year-self.a['national_denominator']['year'])
            self.guarded(s,f'Y{row}',f"'Assumptions'!$B$25*(1+'Assumptions'!$B$27)^(X{row}-'Assumptions'!$B$24)",["'Assumptions'!$B$25","'Assumptions'!$B$27",f'X{row}',"'Assumptions'!$B$24"],den)
            self.guarded(s,f'Z{row}',f'K{row}/Y{row}',[f'K{row}',f'Y{row}'],out['facility_electricity_twh']/den,PCT)
            self.put(s,f'AA{row}',cfg['anchor_source_id'])
            pr,pv = self.pln_forecast_rows[year]
            self.link(s,f'AB{row}',f"'Assumptions'!B{pr}",pv)
            self.guarded(s,f'AC{row}',f'K{row}/AB{row}',[f'K{row}',f'AB{row}'],out['facility_electricity_twh']/pv,PCT)
            self.put(s,f'AD{row}','ELEC003')
            self.outlook_rows[(name,year)] = row
            row += 1
        self.size_model(s)
        self.wb[s].column_dimensions['W'].width = 29
        for col in ['X','Y','Z','AA','AB','AC','AD']:
            self.wb[s].column_dimensions[col].width = 24
        self.wb[s].auto_filter.ref = 'A5:AD14'

    def national(self):
        s = 'National Impact'
        self.title(s,'National electricity impact of the capacity scenarios',
                   'National consumption and PLN sales are separate actual 2025 denominators. Continuous load is average delivered power. Generation capacity at an assumed factor is an energy illustration, not an adequacy forecast.',11)
        self.header(s,5,['Case','IT capacity (MW)','Facility electricity (TWh/year)','National consumption (TWh)','Share of national consumption','PLN sales (TWh)','Share of PLN sales','Continuous delivered power (GW)','Generation capacity at assumed factor (GW)','Denominator year','Source ID'])
        for row, ((mw,case), sr) in enumerate(self.scenario_rows.items(),6):
            out = evaluate(mw,case,self.a)
            self.put(s,f'A{row}',case.title())
            for col, src_col, val in [('B','B',mw),('C','K',out['facility_electricity_twh'])]:
                self.link(s,f'{col}{row}',f"'Scenario Model'!{src_col}{sr}",val)
            self.link(s,f'D{row}',"'Assumptions'!$B$25",self.den['electricity_consumption_twh'])
            self.guarded(s,f'E{row}',f'C{row}/D{row}',[f'C{row}',f'D{row}'],out['facility_electricity_twh']/self.den['electricity_consumption_twh'],PCT)
            self.link(s,f'F{row}',"'Assumptions'!$B$26",self.den['pln_sales_twh'])
            self.guarded(s,f'G{row}',f'C{row}/F{row}',[f'C{row}',f'F{row}'],out['facility_electricity_twh']/self.den['pln_sales_twh'],PCT)
            avg = out['average_facility_load_mw']/1000
            self.guarded(s,f'H{row}',f"C{row}*1000/'Assumptions'!$B$17",[f'C{row}',"'Assumptions'!$B$17"],avg)
            self.guarded(s,f'I{row}',f"H{row}/'Assumptions'!$B$21",[f'H{row}',"'Assumptions'!$B$21"],avg/self.a['generation_capacity_factor'])
            self.link(s,f'J{row}',"'Assumptions'!$B$24",self.a['national_denominator']['year'],'0')
            self.put(s,f'K{row}',self.den['source_id'])
        for col in range(1,12):
            self.wb[s].column_dimensions[get_column_letter(col)].width = 26 if col > 2 else 17
        self.wb[s].auto_filter.ref = 'A5:K23'
        self.wb[s].print_area = 'A1:K23'

    def executive(self):
        s = 'Executive Summary'
        self.title(s,'Indonesia AI and data center resource model',
                   f'Research cutoff {self.a["research_cutoff"]}. Annual steady-state scenarios. National metered data center electricity and water consumption are unavailable in the assembled public sources.',6)
        self.header(s,5,['Additional 1 GW IT','Low','Base','High','Unit','Definition'])
        metric_rows = [
            (6,'IT nameplate capacity','B',1000,'GW',1000),
            (7,'Facility electricity','K',1,'TWh/year',None),
            (8,'Average facility load','L',1000,'GW',None),
            (9,'Full IT load facility power','M',1000,'GW',None),
            (10,'Direct cooling consumption','N',1e6,'million m³/year',None),
            (11,'Direct cooling consumption','O',1,'m³/day',None),
            (12,'Annual electricity cost','P',1e12,'Rp trillion/year',None),
            (13,'Annual water cost','Q',1e9,'Rp billion/year',None),
            (14,'Annual resource cost','R',1e12,'Rp trillion/year',None),
            (15,'Water cost share','T',1,'%',None),
        ]
        for r,label,col,div,unit,_ in metric_rows:
            self.put(s,f'A{r}',label)
            self.put(s,f'E{r}',unit)
            for c,case in enumerate(CASE_COL,2):
                sr = self.scenario_rows[(1000,case)]
                cache = self.caches['Scenario Model'][f'{col}{sr}']/div
                self.guarded(s,f'{get_column_letter(c)}{r}',f"'Scenario Model'!{col}{sr}/{div:g}",[f"'Scenario Model'!{col}{sr}"],cache,PCT if r == 15 else NUM)
        for r,label,col in [(16,'Share of national consumption','E'),(17,'Share of PLN sales','G')]:
            self.put(s,f'A{r}',label)
            self.put(s,f'E{r}','% of actual 2025')
            for c,case in enumerate(CASE_COL,2):
                nr = self.scenario_rows[(1000,case)]
                self.link(s,f'{get_column_letter(c)}{r}',f"'National Impact'!{col}{nr}",self.caches['National Impact'][f'{col}{nr}'],PCT)
        for r,note in [(19,'Water is modeled direct on-site consumptive cooling water per IT kWh. Observed site withdrawals, domestic water and electricity-supply water are excluded.'),(20,'The low case dry-cooling boundary allows zero evaporative cooling consumption. It does not mean zero total site water use.'),(21,'Energy cost uses the dated PLN I-4 high-voltage schedule. Water cost uses the Jakarta large-industrial marginal tariff proxy. Taxes, demand charges and connection costs are excluded.'),(22,'PUE and WUE combinations are conditional stresses, not measured fleet averages or validated joint cooling designs.')]:
            self.wb[s].merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
            self.put(s,f'A{r}',note).alignment = Alignment(wrap_text=True,vertical='center')
            self.wb[s].row_dimensions[r].height = 34
        self.header(s,24,['Dated Greater Jakarta estimate','Low','Base','High','Unit','Definition'])
        for r,label,col,div,unit in [(25,'Operational IT market estimate','B',1,'MW'),(26,'Modeled facility electricity','K',1,'TWh/year'),(27,'Modeled direct water','N',1e6,'million m³/year'),(28,'Modeled resource cost','R',1e12,'Rp trillion/year')]:
            self.put(s,f'A{r}',label)
            self.put(s,f'E{r}',unit)
            for c,case in enumerate(CASE_COL,2):
                sr = 30+list(CASE_COL).index(case)
                self.guarded(s,f'{get_column_letter(c)}{r}',f"'Scenario Model'!{col}{sr}/{div:g}",[f"'Scenario Model'!{col}{sr}"],self.caches['Scenario Model'][f'{col}{sr}']/div)
        self.wb[s].merge_cells('A30:F31')
        self.put(s,'A30',f'{self.a["baseline"]["scope"]} Observation date: {self.a["baseline"]["observation_date"]}. Current PUE scenarios differ from expansion assumptions. These are estimates, not measured consumption.').alignment = Alignment(wrap_text=True,vertical='center')
        self.wb[s].row_dimensions[30].height = 38
        self.header(s,34,['2030 outlook proxy','Conservative','Base Case','Aggressive AI Boom','Unit','Definition'])
        for r,label,col,div,unit in [(35,'IT capacity','B',1,'MW'),(36,'Facility electricity','K',1,'TWh/year'),(37,'Direct cooling consumption','N',1e6,'million m³/year'),(38,'Resource cost','R',1e12,'Rp trillion/year'),(39,'Share of future national consumption','Z',1,'%')]:
            self.put(s,f'A{r}',label)
            self.put(s,f'E{r}',unit)
            for c,name in enumerate(self.a['outlook'],2):
                sr = self.outlook_rows[(name,2030)]
                self.guarded(s,f'{get_column_letter(c)}{r}',f"'2030 Outlook'!{col}{sr}/{div:g}",[f"'2030 Outlook'!{col}{sr}"],self.caches['2030 Outlook'][f'{col}{sr}']/div,PCT if r == 39 else NUM)
        self.wb[s].merge_cells('A41:F42')
        self.put(s,'A41','Outlook combines Greater Jakarta with separately sourced operating Surabaya, Batam construction, and Batam/Bintan planned capacity. It remains an incomplete nationwide inventory. The current metro estimate and partial facility inventory overlap and must not be added.').alignment = Alignment(wrap_text=True,vertical='center')
        self.wb[s].row_dimensions[41].height = 38
        self.put(s,'F6','Expansion cases').alignment = Alignment(wrap_text=True)
        self.put(s,'F25','ELEC004 H2 2025').alignment = Alignment(wrap_text=True)
        self.put(s,'F35','See outlook source IDs').alignment = Alignment(wrap_text=True)
        for col,width in {'A':42,'B':22,'C':22,'D':25,'E':25,'F':33}.items():
            self.wb[s].column_dimensions[col].width = width
        self.wb[s].freeze_panes = 'B6'
        self.wb[s].print_area = 'A1:F42'
        for row in [7,10,14,26,36,38]:
            for cell in self.wb[s][row][:6]:
                cell.fill = PatternFill('solid',fgColor='EBF3F5')
                cell.border = Border(bottom=Side(style='thin',color='B7CDD2'))
        for row in self.wb[s].iter_rows():
            for cell in row:
                if cell.data_type == 'f':
                    cell.font = Font(name='Arial', size=10, color=NAVY)

    def build(self):
        for sheet,name in RAW_SHEETS.items():
            self.raw_table(sheet,read_dataset(name))
        self.raw_table('Sources',pd.read_csv(ROOT/'sources.csv'))
        electricity = self.tables['Indonesia Electricity']
        match = electricity[electricity.year.eq(self.a['national_denominator']['year'])]
        if len(match) != 1:
            raise ValueError('Expected exactly one dated national denominator record')
        self.den = match.iloc[0]
        self.den_row = int(match.index[0])+2
        p = self.tables['Data Center Projects']
        mask = p['include_in_inventory'].astype(str).str.lower().eq('true') & p['capacity_basis'].astype(str).str.upper().eq('IT')
        self.inventory_mw = float(pd.to_numeric(p.loc[mask,'operational_mw'],errors='coerce').sum(min_count=1))
        # Connection estimates recalculate from their own apparent-power inputs.
        # Reported real MW remains an original observation, never replaced.
        mva_col = get_column_letter(p.columns.get_loc('power_connection_mva') + 1)
        mw_col = get_column_letter(p.columns.get_loc('power_connection_mw') + 1)
        pf_col = get_column_letter(p.columns.get_loc('conversion_power_factor') + 1)
        for r, (_, record) in enumerate(p.iterrows(), 2):
            if record.get('connection_data_class') == 'E':
                self.guarded('Data Center Projects', f'{mw_col}{r}',
                             f"{mva_col}{r}*'Assumptions'!$B$20",
                             [f'{mva_col}{r}', "'Assumptions'!$B$20"], record['power_connection_mw'])
                self.link('Data Center Projects', f'{pf_col}{r}', "'Assumptions'!$B$20", self.a['power_factor']['base'])
        # A zero contribution means excluded or undisclosed, never zero observed
        # facility capacity. Separate row guards preserve the raw capacity cells.
        cc = get_column_letter(len(p.columns)+1)
        ws = self.wb['Data Center Projects']
        cell = self.put('Data Center Projects',f'{cc}1','disclosed_inventory_contribution_it_mw (E)')
        cell.font = Font(name='Arial',size=10,bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid',fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True,vertical='center')
        cols = {name:get_column_letter(p.columns.get_loc(name)+1) for name in ['operational_mw','include_in_inventory','capacity_basis']}
        for r,(_,record) in enumerate(p.iterrows(),2):
            cap = record['operational_mw']
            included = str(record['include_in_inventory']).lower() == 'true' and str(record['capacity_basis']).upper() == 'IT' and pd.notna(cap)
            self.formula('Data Center Projects',f'{cc}{r}',f'=IF(AND({cols["include_in_inventory"]}{r}=TRUE,{cols["capacity_basis"]}{r}="IT",ISNUMBER({cols["operational_mw"]}{r})),{cols["operational_mw"]}{r},0)',float(cap) if included else 0)
        ws.column_dimensions[cc].width = 29
        ws.auto_filter.ref = ws.dimensions
        ws.print_area = ws.dimensions
        self.assumptions()
        self.model()
        self.sensitivity()
        self.outlook()
        self.national()
        self.executive()
        for sheet in SHEETS:
            self.fit_wrapped_rows(sheet)
        self.wb.active = 0
        self.wb['Executive Summary'].sheet_view.selection[0].activeCell = 'A1'
        return self


def _write_formula_caches(path, caches):
    """Insert standards-compliant <v> caches without replacing formula text."""
    ns = {'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    ET.register_namespace('',ns['m'])
    with ZipFile(path) as z:
        entries = {info.filename:(info,z.read(info.filename)) for info in z.infolist()}
    for index,sheet in enumerate(SHEETS,1):
        key = f'xl/worksheets/sheet{index}.xml'
        info,data = entries[key]
        tree = ET.fromstring(data)
        for cell in tree.findall('.//m:c',ns):
            address = cell.get('r')
            if address not in caches[sheet]:
                continue
            value = caches[sheet][address]
            if isinstance(value,float) and not math.isfinite(value):
                raise ValueError(f'Non-finite formula cache: {sheet}!{address}')
            v = cell.find('m:v',ns)
            if v is None:
                v = ET.SubElement(cell,f'{{{ns["m"]}}}v')
            if value is None:
                cell.set('t','str')
                v.text = 'n.a.'
            elif isinstance(value,str):
                cell.set('t','str')
                v.text = value
            else:
                cell.set('t','n')
                v.text = repr(value)
        entries[key] = (info,ET.tostring(tree,encoding='utf-8',xml_declaration=True))
    temp = path.with_suffix('.xlsx.tmp')
    with ZipFile(temp,'w',ZIP_DEFLATED) as z:
        for key,(info,data) in entries.items():
            z.writestr(info,data)
    temp.replace(path)


def validate_saved_workbook(path,builder):
    formulas = load_workbook(path,data_only=False)
    values = load_workbook(path,data_only=True)
    if formulas.sheetnames != SHEETS:
        raise AssertionError('Requested worksheet names/order changed')
    errors=[]
    for sheet,cells in builder.caches.items():
        for address,expected in cells.items():
            actual = values[sheet][address].value
            if isinstance(expected,(int,float)):
                if not isinstance(actual,(int,float)) or not math.isclose(actual,expected,rel_tol=1e-11,abs_tol=1e-7):
                    errors.append(f'{sheet}!{address}: cache mismatch')
            elif actual != (expected if expected is not None else 'n.a.'):
                errors.append(f'{sheet}!{address}: nonnumeric cache mismatch')
            formula = formulas[sheet][address].value
            if not isinstance(formula,str) or not formula.startswith('=') or '#REF!' in formula:
                errors.append(f'{sheet}!{address}: invalid formula')
    # Reconcile all 18 base/low/high outputs with the pipeline, not only headline values.
    frame = read_dataset('capacity_scenarios')
    for _,record in frame.iterrows():
        r = builder.scenario_rows[(record.it_capacity_mw,record['case'])]
        for col,key in [('I','it_electricity_mwh'),('J','facility_electricity_mwh'),('N','water_m3_per_year'),('P','electricity_cost_rp'),('Q','water_cost_rp'),('R','total_resource_cost_rp')]:
            if not math.isclose(values['Scenario Model'][f'{col}{r}'].value,float(record[key]),rel_tol=1e-10,abs_tol=1e-5):
                errors.append(f'Pipeline reconciliation: {col}{r}/{key}')
    result = {
        'workbook':path.name,'worksheet_names':formulas.sheetnames,
        'formula_cache_method':'Python independently evaluates model inputs and inserts Open XML cached values. Excel formulas retained; automatic full recalculation requested on open.',
        'formula_cache_reopen_check':not errors,'pipeline_capacity_reconciliation':not errors,
        'raw_rows':{sheet:len(frame) for sheet,frame in builder.tables.items()},
        'partial_disclosed_operational_it_mw':builder.inventory_mw,
        'artifact_tool_recalculation':'Pending independent QA; run outputs/.qa/verify_workbook.mjs when the bundled runtime is available.',
        'native_excel_recalculation':'Not performed. No native Excel engine available in this task.',
        'visual_review':'Pending all-sheet rendered QA.','errors':errors,
    }
    if errors:
        raise AssertionError('; '.join(errors[:12]))
    return result


def export_excel(output_path=None):
    """Export one workbook; all processed datasets must already exist."""
    path = Path(output_path) if output_path else ROOT/'outputs'/'indonesia_ai_datacenter_model.xlsx'
    path.parent.mkdir(parents=True,exist_ok=True)
    builder = Builder(assumptions()).build()
    builder.wb.save(path)
    _write_formula_caches(path,builder.caches)
    result = validate_saved_workbook(path,builder)
    write_json(ROOT/'outputs'/'workbook_validation.json',result)
    print(f'Exported {path.name}: all 15 requested sheets; formula caches reconcile with the Python capacity scenarios.')
    return path


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=None)
    args = parser.parse_args()
    export_excel(args.output)
